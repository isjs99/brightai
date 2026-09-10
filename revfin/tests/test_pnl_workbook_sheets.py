import json

import pytest
from openpyxl import load_workbook

from revfin.fixtures import load_fixtures
from revfin.pnl import ANALYTICS, build_pnl
from revfin.sheets import push_workbook, resolve_target
from revfin.workbook import build_workbook, write_xlsx
from tests.conftest import utc


@pytest.fixture
def model(settings, db, home):
    load_fixtures(settings, db, home / "fixtures")
    return build_pnl(settings, db, now=utc(2026, 9, 10, 12))


def test_pnl_lines_reconcile_to_summary_figures(model):
    uk = next(e for e in model.entities if e.slug == "uk")
    assert model.months == ["2026-06", "2026-07", "2026-08", "2026-09"]
    # Same August figures the markdown summary reports.
    assert uk.revenue("2026-08") == 22642.0
    assert uk.cost("cost_of_sales", "2026-08") == 3900.0        # contractor 1,500 + ads 2,400
    assert uk.gross_profit("2026-08") == 18742.0
    assert uk.net_result("2026-08") == 3002.01
    # fx / internal / intercompany never land in a P&L line.
    assert {l.key for l in uk.lines}.isdisjoint({"fx", "internal", "intercompany"})
    assert all(r["pnl_section"] == "excluded" for r in model.ledger if r["category"] in ("fx", "internal", "intercompany"))
    # J Patel's 350 inflow sits in uncategorised as a negative cost, not revenue.
    unc = next(l for l in uk.lines if l.key == "uncategorised")
    assert unc.values["2026-08"] == pytest.approx(350.0 - 6565.0)
    assert uk.clients["Acme Beauty Ltd"]["2026-08"] == 15400.0
    assert uk.client_meta["Acme Beauty Ltd"]["receipts"] == 5
    assert list(uk.vendors)[0] == "A Smith"  # 4 months of payroll beats the one-off 6,500


def test_group_consolidates_in_gbp_and_nets_intercompany(model):
    g = model.group
    uk = next(e for e in model.entities if e.slug == "uk")
    de = next(e for e in model.entities if e.slug == "de")
    assert g.currency == "GBP"
    assert g.revenue("2026-08") == pytest.approx(uk.revenue("2026-08") + round(de.revenue("2026-08") * 0.852, 2), abs=0.02)
    # The 2,000 EUR moved UK -> DE must not appear on either side.
    assert "Brightform Social Media Limited" not in de.clients
    assert g.cash_total["2026-09"] == pytest.approx(uk.cash_total["2026-09"] + round(de.cash_total["2026-09"] * 0.852, 2), abs=0.02)
    assert any(n.startswith("group: EUR->GBP 0.8520") for n in model.fx_notes)


def test_analytics_metrics(model):
    uk = next(e for e in model.entities if e.slug == "uk")
    a = uk.analytics
    assert set(a) == {key for key, _, _ in ANALYTICS}
    assert a["revenue_mom_pct"]["2026-06"] is None
    assert a["revenue_mom_pct"]["2026-08"] == pytest.approx((22642 - 19242) / 19242, abs=1e-4)
    assert a["revenue_ytd"]["2026-08"] == pytest.approx(19242 + 19242 + 22642)
    assert a["revenue_t3m"]["2026-08"] == pytest.approx((19242 + 19242 + 22642) / 3, abs=0.01)
    assert a["gross_margin_pct"]["2026-08"] == pytest.approx(18742 / 22642, abs=1e-4)
    assert a["active_clients"]["2026-08"] == 2
    assert a["top_client_share_pct"]["2026-08"] == pytest.approx(15400 / 22642, abs=1e-4)
    assert a["cumulative_net"]["2026-08"] == pytest.approx(6762.01 + 6917.01 + 3002.01)
    assert a["net_burn"]["2026-08"] == 0.0 and a["runway_months"]["2026-08"] is None
    assert a["cash_end"]["2026-09"] == uk.cash_total["2026-09"]
    assert a["uncategorised_spend"]["2026-08"] == 6215.0


def test_month_range_can_be_pinned(settings, db, home):
    load_fixtures(settings, db, home / "fixtures")
    model = build_pnl(settings, db, "2026-01", "2026-03", now=utc(2026, 9, 10))
    assert model.months == ["2026-01", "2026-02", "2026-03"]
    assert model.group.revenue("2026-02") == 0.0


def test_workbook_tabs_and_xlsx(model, settings, tmp_path):
    tabs = build_workbook(model, settings.account_nicknames)
    assert [t.name for t in tabs] == ["Overview", "P&L Group", "P&L UK", "P&L DE", "Analytics", "Clients", "Vendors",
                                      "Cash", "Uncategorised", "Ledger", "Mapping"]
    pnl_uk = next(t for t in tabs if t.name == "P&L UK")
    assert pnl_uk.rows[0] == ["Line", "Jun 2026", "Jul 2026", "Aug 2026", "Sep 2026", "Total"]
    revenue_total = next(r for r in pnl_uk.rows if r[0] == "Total revenue")
    assert revenue_total[3] == 22642.0 and revenue_total[-1] == pytest.approx(19242 + 19242 + 22642 + 12000)
    net = next(r for r in pnl_uk.rows if r[0] == "Net cash result")
    assert net[3] == 3002.01
    assert len(pnl_uk.charts) == 2
    ledger = next(t for t in tabs if t.name == "Ledger")
    assert len(ledger.rows) == len(model.ledger) + 1
    # every row has the same width as its header, so Sheets/xlsx columns line up
    for tab in tabs:
        width = len(tab.rows[0])
        assert all(len(r) == width for r in tab.rows), tab.name

    path = write_xlsx(tabs, tmp_path / "pnl.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == [t.name for t in tabs]
    ws = wb["P&L UK"]
    assert ws.freeze_panes == "B2"
    assert len(ws._charts) == 2
    cell = ws.cell(row=revenue_total_row(pnl_uk) + 1, column=4)
    assert cell.value == 22642.0 and cell.number_format.startswith("#,##0.00")
    assert wb["Cash"]._charts and wb["Ledger"].max_row == len(model.ledger) + 1


def revenue_total_row(tab):
    return next(i for i, r in enumerate(tab.rows) if r[0] == "Total revenue")


class FakeSheetsApi:
    """Records the calls push_workbook makes and simulates sheet creation."""

    def __init__(self, existing_titles=("Sheet1",)):
        self.sheets = [{"properties": {"title": t, "sheetId": i}, "charts": [{"chartId": 900 + i}] if t == "Cash" else []}
                       for i, t in enumerate(existing_titles)]
        self.batches: list[list[dict]] = []
        self.cleared: list[str] = []
        self.updates: list[tuple[str, list]] = []

    def get(self):
        return {"properties": {"title": "Brightform P&L (revfin)"}, "sheets": [dict(s) for s in self.sheets]}

    def batch_update(self, requests):
        self.batches.append(requests)
        for req in requests:
            if "addSheet" in req:
                self.sheets.append({"properties": {"title": req["addSheet"]["properties"]["title"], "sheetId": 100 + len(self.sheets)}, "charts": []})
            if "deleteSheet" in req:
                self.sheets = [s for s in self.sheets if s["properties"]["sheetId"] != req["deleteSheet"]["sheetId"]]
        return {}

    def values_clear(self, range_):
        self.cleared.append(range_)

    def values_update(self, range_, values):
        self.updates.append((range_, values))


def test_push_workbook_creates_tabs_writes_values_and_charts(model, settings):
    tabs = build_workbook(model, settings.account_nicknames)
    api = FakeSheetsApi(existing_titles=("Sheet1", "Cash"))
    stats = push_workbook(api, tabs)
    assert stats["tabs"] == len(tabs) and stats["charts"] == 2 * 3 + 1
    titles = [s["properties"]["title"] for s in api.sheets]
    assert "Sheet1" not in titles and set(titles) == {t.name for t in tabs}
    # Cash tab existed with a chart: it must be deleted before the new one is added.
    flat = [req for batch in api.batches for req in batch]
    assert {"deleteEmbeddedObject": {"objectId": 901}} in flat
    assert sum(1 for r in flat if "addChart" in r) == stats["charts"]
    assert "'P&L UK'!A1:ZZ" in api.cleared
    ledger_update = next(v for rng, v in api.updates if rng == "'Ledger'!A1")
    assert ledger_update[0][0] == "Date" and len(ledger_update) == len(model.ledger) + 1
    assert all("" == c or not isinstance(c, float) or c == c for row in ledger_update for c in row)
    # None becomes "" for the API, never the string "None".
    assert not any(c == "None" for rng, v in api.updates for row in v for c in row)
    chart = next(r for r in flat if "addChart" in r)["addChart"]["chart"]["spec"]["basicChart"]
    assert chart["headerCount"] == 1 and chart["chartType"] in ("COLUMN", "LINE")


def test_resolve_target_accepts_url_and_requires_key(settings, monkeypatch, home):
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_FILE", "secrets/google-service-account.json")
    sid, key = resolve_target(settings, "https://docs.google.com/spreadsheets/d/1abcDEF/edit#gid=0")
    assert sid == "1abcDEF" and key == (home / "secrets/google-service-account.json").resolve()
    monkeypatch.delenv("GOOGLE_SERVICE_ACCOUNT_FILE")
    monkeypatch.delenv("REVFIN_SHEET_ID", raising=False)
    from revfin.config import ConfigError
    with pytest.raises(ConfigError, match="REVFIN_SHEET_ID"):
        resolve_target(settings, None)


def test_pnl_cli_writes_workbook_and_json(home, tmp_path):
    from typer.testing import CliRunner
    from revfin import cli
    runner = CliRunner()
    assert runner.invoke(cli.app, ["load-fixtures"], catch_exceptions=False).exit_code == 0
    out = tmp_path / "pnl.xlsx"
    result = runner.invoke(cli.app, ["pnl", "--xlsx", str(out), "--json", str(tmp_path / "pnl.json")], catch_exceptions=False)
    assert result.exit_code == 0, result.output
    assert "Group consolidated (GBP)" in result.output and out.exists()
    fig = json.loads((tmp_path / "pnl.json").read_text())
    assert fig["group"]["derived"]["2026-08"]["net_result"] > 0
    result = runner.invoke(cli.app, ["sheets", "push"], catch_exceptions=False)
    assert result.exit_code == 1 and "REVFIN_SHEET_ID" in result.output
