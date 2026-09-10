"""Turn a PnlModel into tabs, and render tabs to .xlsx. sheets.py renders the
same tabs to Google Sheets so both outputs are always identical."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from .pnl import ANALYTICS, COST_SECTIONS, SECTION_LABELS, EntityPnl, PnlModel
from .util import month_label

NUMBER_FORMAT = "#,##0.00;[Red]-#,##0.00"
PERCENT_FORMAT = "0.0%"
INT_FORMAT = "0"


@dataclass
class Chart:
    kind: str  # "column" | "line"
    title: str
    header_row: int          # 0-based row holding the month labels (categories)
    series_rows: list[int]   # 0-based rows, first cell is the series name
    first_col: int           # 0-based column of the series label
    last_col: int            # 0-based inclusive last data column
    anchor_row: int
    anchor_col: int


@dataclass
class Tab:
    name: str
    rows: list[list]
    freeze_rows: int = 1
    freeze_cols: int = 1
    bold_rows: set[int] = field(default_factory=set)
    percent_rows: set[int] = field(default_factory=set)
    int_rows: set[int] = field(default_factory=set)
    charts: list[Chart] = field(default_factory=list)
    col_widths: dict[int, int] = field(default_factory=dict)
    note: str | None = None


def _label(m: str) -> str:
    return month_label(m)[:3] + " " + m[:4]


# -- tab builders ------------------------------------------------------------------

def pnl_tab(v: EntityPnl, title: str) -> Tab:
    months = v.months
    header = ["Line", *[_label(m) for m in months], "Total"]
    rows: list[list] = [header]
    bold: set[int] = {0}
    pct: set[int] = set()
    chart_rows: dict[str, int] = {}

    def add(label: str, values: list[float | None], is_bold=False, is_pct=False, key: str | None = None):
        total = None if is_pct else round(sum(x for x in values if x is not None), 2)
        rows.append([label, *values, total])
        if is_bold:
            bold.add(len(rows) - 1)
        if is_pct:
            pct.add(len(rows) - 1)
        if key:
            chart_rows[key] = len(rows) - 1

    for section in ("revenue", "cost_of_sales", "opex", "tax", "equity"):
        rows.append([SECTION_LABELS[section]] + [None] * (len(months) + 1))
        bold.add(len(rows) - 1)
        sign = 1 if section == "revenue" else -1
        for line in v.lines:
            if line.section != section:
                continue
            add(f"  {line.label}", [round(sign * line.values.get(m, 0.0), 2) for m in months])
        section_values = [round(sign * v.sections[section].get(m, 0.0), 2) for m in months]
        add(f"Total {SECTION_LABELS[section].lower()}", section_values, is_bold=True, key=section)
        if section == "cost_of_sales":
            add("Gross profit", [v.gross_profit(m) for m in months], is_bold=True, key="gross_profit")
            add("Gross margin %", [_safe_pct(v.gross_profit(m), v.revenue(m)) for m in months], is_pct=True)
            rows.append([""] + [None] * (len(months) + 1))
        elif section == "opex":
            add("Operating result", [v.operating_result(m) for m in months], is_bold=True, key="operating")
            rows.append([""] + [None] * (len(months) + 1))
        elif section == "tax":
            add("Net cash result", [v.net_result(m) for m in months], is_bold=True, key="net")
            add("Net margin %", [_safe_pct(v.net_result(m), v.revenue(m)) for m in months], is_pct=True)
            cumulative, cum = [], 0.0
            for m in months:
                cum = round(cum + v.net_result(m), 2)
                cumulative.append(cum)
            rows.append(["Cumulative net", *cumulative, None])
            chart_rows["cumulative"] = len(rows) - 1
            rows.append([""] + [None] * (len(months) + 1))
        elif section == "equity":
            add("Net after owner drawings", [v.net_after_equity(m) for m in months], is_bold=True)
            rows.append([""] + [None] * (len(months) + 1))

    total_costs = [round(v.cost("cost_of_sales", m) + v.cost("opex", m) + v.cost("tax", m), 2) for m in months]
    rows.append(["Total costs (for chart)", *total_costs, round(sum(total_costs), 2)])
    chart_rows["costs"] = len(rows) - 1
    rows.append(["Cash at month end", *[v.cash_total.get(m) for m in months], None])
    chart_rows["cash"] = len(rows) - 1

    last_col = len(months)
    anchor_col = last_col + 3
    charts = [
        Chart("column", f"{title}: revenue, costs, net by month", 0,
              [chart_rows["revenue"], chart_rows["costs"], chart_rows["net"]], 0, last_col, 1, anchor_col),
        Chart("line", f"{title}: cumulative net and cash", 0,
              [chart_rows["cumulative"], chart_rows["cash"]], 0, last_col, 20, anchor_col),
    ]
    widths = {0: 34, **{i + 1: 13 for i in range(len(months) + 1)}}
    return Tab(f"P&L {title}", rows, bold_rows=bold, percent_rows=pct, charts=charts, col_widths=widths,
               note=f"{v.name}. Cash basis, {v.currency}. Costs shown as positive numbers.")


def analytics_tab(model: PnlModel) -> Tab:
    months = model.months
    rows: list[list] = [["Metric", *[_label(m) for m in months]]]
    bold, pct, ints = {0}, set(), set()
    for v in model.all_views():
        rows.append([f"{v.name} ({v.currency})"] + [None] * len(months))
        bold.add(len(rows) - 1)
        for key, label, kind in ANALYTICS:
            rows.append([f"  {label}", *[v.analytics[key].get(m) for m in months]])
            if kind == "pct":
                pct.add(len(rows) - 1)
            elif kind == "count":
                ints.add(len(rows) - 1)
        rows.append([""] + [None] * len(months))
    return Tab("Analytics", rows, bold_rows=bold, percent_rows=pct, int_rows=ints,
               col_widths={0: 40, **{i + 1: 13 for i in range(len(months))}},
               note="Runway = cash at month end / average net burn over the trailing 3 months.")


def clients_tab(model: PnlModel) -> Tab:
    months = model.months
    rows: list[list] = [["Entity", "Client", "Currency", *[_label(m) for m in months], "Total", "Share of entity revenue",
                         "Receipts", "First receipt", "Last receipt"]]
    pct = set()
    for v in model.entities:
        entity_total = sum(sum(vals.values()) for vals in v.clients.values())
        for client, vals in v.clients.items():
            total = round(sum(vals.values()), 2)
            meta = v.client_meta.get(client, {})
            rows.append([v.slug, client, v.currency, *[vals.get(m, 0.0) for m in months], total,
                         _safe_pct(total, entity_total), meta.get("receipts"), meta.get("first"), meta.get("last")])
    return Tab("Clients", rows, freeze_cols=2, bold_rows={0}, int_rows=set(), percent_rows=set(),
               col_widths={0: 8, 1: 34, 2: 9, **{i + 3: 12 for i in range(len(months) + 1)}, len(months) + 4: 12},
               note="Only legs tagged as revenue. Share column is a fraction of that entity's revenue over the whole range.")


def vendors_tab(model: PnlModel) -> Tab:
    months = model.months
    rows: list[list] = [["Entity", "Vendor / payee", "Category", "Currency", *[_label(m) for m in months], "Total", "Items"]]
    for v in model.entities:
        for vendor, d in v.vendors.items():
            rows.append([v.slug, vendor, d["category"], v.currency, *[d["months"].get(m, 0.0) for m in months], d["total"], d["items"]])
    return Tab("Vendors", rows, freeze_cols=2, bold_rows={0},
               col_widths={0: 8, 1: 34, 2: 14, 3: 9, **{i + 4: 12 for i in range(len(months) + 1)}},
               note="Outflows in cost of sales, opex and tax, sorted by total within each entity.")


def cash_tab(model: PnlModel) -> Tab:
    months = model.months
    rows: list[list] = [["Entity", "Account", "Currency", *[_label(m) for m in months]]]
    bold = {0}
    chart_rows = []
    for v in model.entities:
        for label, series in v.cash.items():
            rows.append([v.slug, label, v.cash_currency.get(label, ""), *[series.get(m) for m in months]])
        rows.append([v.slug, f"Total in {v.currency}", v.currency, *[v.cash_total.get(m) for m in months]])
        bold.add(len(rows) - 1)
        chart_rows.append(len(rows) - 1)
    rows.append(["group", f"Group total in {model.group_currency}", model.group_currency, *[model.group.cash_total.get(m) for m in months]])
    bold.add(len(rows) - 1)
    chart_rows.append(len(rows) - 1)
    last_col = len(months) + 2
    charts = [Chart("line", "Cash at month end", 0, chart_rows, 1, last_col, len(rows) + 2, 0)]
    return Tab("Cash", rows, freeze_cols=3, bold_rows=bold, charts=charts,
               col_widths={0: 8, 1: 40, 2: 9, **{i + 3: 13 for i in range(len(months))}},
               note="Month-end balance per account: the newer of the last completed leg's balance and the last sync snapshot.")


LEDGER_COLUMNS = [
    ("date", "Date"), ("month", "Month"), ("entity", "Entity"), ("account_label", "Account"), ("type", "Type"),
    ("state", "State"), ("who", "Counterparty / merchant"), ("reference", "Reference"), ("category", "Category"),
    ("pnl_section", "P&L section"), ("pnl_line", "P&L line"), ("amount", "Amount"), ("currency", "Currency"),
    ("amount_base", "In entity currency"), ("base_currency", "Entity currency"), ("amount_group", "In group currency"),
    ("group_currency", "Group currency"), ("counted", "In P&L"), ("balance_after", "Balance after"),
    ("transaction_id", "Transaction id"), ("leg_index", "Leg"), ("rule_id", "Rule"),
]


def ledger_tab(model: PnlModel, nicknames: dict[str, str]) -> Tab:
    rows: list[list] = [[label for _, label in LEDGER_COLUMNS]]
    for r in model.ledger:
        r = dict(r)
        r["date"] = (r["created_at"] or "")[:10]
        r["account_label"] = nicknames.get(r["account_id"], r.get("account_name") or r["account_id"])
        r["counted"] = "yes" if r["counted"] else "no"
        rows.append([r.get(key) for key, _ in LEDGER_COLUMNS])
    ints = set()
    return Tab("Ledger", rows, freeze_cols=0, bold_rows={0}, int_rows=ints,
               col_widths={0: 11, 1: 8, 2: 7, 3: 18, 4: 12, 5: 10, 6: 30, 7: 30, 8: 14, 9: 13, 10: 26, 11: 12, 12: 8, 13: 14, 14: 8, 15: 14, 16: 8, 17: 7, 18: 13, 19: 16, 20: 5, 21: 16},
               note="Every transaction leg in the database. Pivot on this for anything the other tabs do not show.")


def uncategorised_tab(model: PnlModel) -> Tab:
    rows: list[list] = [["Date", "Entity", "Counterparty / merchant", "Reference", "Type", "Amount", "Currency", "Transaction id"]]
    for r in model.uncategorised:
        rows.append([r["created_at"][:10], r["entity"], r["who"], r.get("reference"), r["type"], r["amount"], r["currency"], r["transaction_id"]])
    return Tab("Uncategorised", rows, freeze_cols=0, bold_rows={0},
               col_widths={0: 11, 1: 7, 2: 34, 3: 34, 4: 12, 5: 12, 6: 8, 7: 16},
               note="Add a rule per row in config.yaml under categories, then run `revfin categorise` and push again.")


def mapping_tab(model: PnlModel) -> Tab:
    rows: list[list] = [["Category", "P&L section", "P&L line"]]
    for line in model.lines_map:
        rows.append([line.category, SECTION_LABELS[line.section], line.label])
    for cat in ("fx", "internal", "intercompany"):
        rows.append([cat, "Excluded", "Not income or spend (currency exchange, own-account and intercompany transfers)"])
    return Tab("Mapping", rows, freeze_cols=0, bold_rows={0}, col_widths={0: 18, 1: 20, 2: 60},
               note="From config.yaml. Change the pnl.lines block there to restructure the P&L.")


def overview_tab(model: PnlModel) -> Tab:
    months = model.months
    latest = months[-1]
    prev = months[-2] if len(months) > 1 else None
    rows: list[list] = [
        ["Brightform P&L (revfin)", None, None, None, None],
        [f"Generated {model.generated_at.strftime('%Y-%m-%d %H:%M')} UTC. Cash basis from Revolut Business. Months {months[0]} to {months[-1]}.", None, None, None, None],
        ["FX: " + ("; ".join(model.fx_notes) if model.fx_notes else "no conversions needed"), None, None, None, None],
        [None] * 5,
        ["View", "Metric", _label(latest), f"{_label(prev)}" if prev else "Prior month", "Change"],
    ]
    bold = {0, 4}
    pct = set()
    keys = [("revenue", "Revenue"), ("gross_profit", "Gross profit"), ("opex", "Operating expenses"), ("net_result", "Net cash result"),
            ("cash_end", "Cash at month end"), ("active_clients", "Paying clients"), ("runway_months", "Runway (months)")]
    for v in model.all_views():
        rows.append([f"{v.name} ({v.currency})", None, None, None, None])
        bold.add(len(rows) - 1)
        for key, label in keys:
            now_v = v.analytics[key].get(latest)
            prev_v = v.analytics[key].get(prev) if prev else None
            change = None if now_v is None or prev_v is None else round(now_v - prev_v, 2)
            rows.append([None, label, now_v, prev_v, change])
        rows.append([None] * 5)
    rows.append(["Tabs", None, None, None, None])
    bold.add(len(rows) - 1)
    for name, what in [
        ("P&L <view>", "monthly cash-basis P&L per entity and consolidated, with charts"),
        ("Analytics", "MoM, trailing 3 month, YTD, margins, concentration, burn, runway"),
        ("Clients", "revenue by client by month"), ("Vendors", "spend by payee by month"),
        ("Cash", "month-end balances per account"), ("Uncategorised", "legs still needing a rule"),
        ("Ledger", "every transaction leg"), ("Mapping", "category to P&L line"),
    ]:
        rows.append([name, what, None, None, None])
    return Tab("Overview", rows, freeze_rows=0, freeze_cols=0, bold_rows=bold, percent_rows=pct,
               col_widths={0: 44, 1: 26, 2: 14, 3: 14, 4: 14})


def build_workbook(model: PnlModel, nicknames: dict[str, str]) -> list[Tab]:
    tabs = [overview_tab(model), pnl_tab(model.group, "Group")]
    tabs += [pnl_tab(v, v.slug.upper()) for v in model.entities]
    tabs += [analytics_tab(model), clients_tab(model), vendors_tab(model), cash_tab(model),
             uncategorised_tab(model), ledger_tab(model, nicknames), mapping_tab(model)]
    return tabs


def _safe_pct(numerator: float | None, denominator: float | None) -> float | None:
    if numerator is None or not denominator:
        return None
    return round(numerator / denominator, 4)


# -- xlsx renderer -----------------------------------------------------------------

def write_xlsx(tabs: list[Tab], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for tab in tabs:
        ws = wb.create_sheet(title=tab.name[:31])
        for r_idx, row in enumerate(tab.rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    if (r_idx - 1) in tab.percent_rows:
                        cell.number_format = PERCENT_FORMAT
                    elif (r_idx - 1) in tab.int_rows or (isinstance(value, int) and (r_idx - 1) in tab.int_rows):
                        cell.number_format = INT_FORMAT
                    else:
                        cell.number_format = NUMBER_FORMAT
                if (r_idx - 1) in tab.bold_rows:
                    cell.font = Font(bold=True)
        if tab.freeze_rows or tab.freeze_cols:
            ws.freeze_panes = ws.cell(row=tab.freeze_rows + 1, column=tab.freeze_cols + 1)
        for col, width in tab.col_widths.items():
            ws.column_dimensions[get_column_letter(col + 1)].width = width
        for chart in tab.charts:
            obj = BarChart() if chart.kind == "column" else LineChart()
            obj.title = chart.title
            obj.height, obj.width = 9, 22
            for row in chart.series_rows:
                data = Reference(ws, min_col=chart.first_col + 1, max_col=chart.last_col + 1, min_row=row + 1, max_row=row + 1)
                obj.add_data(data, from_rows=True, titles_from_data=True)
            cats = Reference(ws, min_col=chart.first_col + 2, max_col=chart.last_col + 1, min_row=chart.header_row + 1, max_row=chart.header_row + 1)
            obj.set_categories(cats)
            ws.add_chart(obj, f"{get_column_letter(chart.anchor_col + 1)}{chart.anchor_row + 1}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
